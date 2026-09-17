from datetime import date, datetime, time
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from planning.application.access import ProjectAccess
from planning.application.errors import InvalidPlanningRequest
from planning.application.interfaces.planning_batch_repository import (
    PlanningBatchExecutor,
)
from planning.application.interfaces.project_management_repositories import (
    ProjectJobsRepository,
)
from planning.application.interfaces.transaction_manager import TransactionManager


class ProjectJobImportInteractor:
    def __init__(
        self,
        access: ProjectAccess,
        repository: ProjectJobsRepository,
        executor: PlanningBatchExecutor,
        transaction_manager: TransactionManager,
    ):
        self._access = access
        self._repository = repository
        self._executor = executor
        self._transaction_manager = transaction_manager

    async def __call__(self, content: bytes, *, apply: bool) -> dict[str, Any]:
        user, project_id = await self._access.dispatcher()
        rows, parse_errors = _parse_workbook(content)
        if not apply:
            return {
                "valid_rows": len(rows),
                "errors": parse_errors,
                "preview": rows[:100],
            }
        result = await self._repository.import_jobs(project_id, rows, user.id)
        result["errors"] = [*parse_errors, *result["errors"]]
        await self._transaction_manager.commit()
        if result["planning_event_id"] is not None:
            self._executor.schedule_project(project_id)
        return result


def _parse_workbook(
    content: bytes,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise InvalidPlanningRequest(
            "The uploaded file is not a valid XLSX workbook",
            code="INVALID_IMPORT_FILE",
        ) from error
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    header = next(iterator, None)
    if not header:
        raise InvalidPlanningRequest(
            "The workbook is empty", code="INVALID_IMPORT_FILE"
        )
    keys = [str(value or "").strip().casefold() for value in header]
    required = {"address", "sla_date", "work_type"}
    if not required.issubset(keys):
        raise InvalidPlanningRequest(
            "Required columns: address, sla_date, work_type",
            code="INVALID_IMPORT_COLUMNS",
        )
    result: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for number, cells in enumerate(iterator, start=2):
        raw = dict(zip(keys, cells, strict=False))
        if not any(value is not None and value != "" for value in cells):
            continue
        try:
            address = str(raw.get("address") or "").strip()
            work_type = str(raw.get("work_type") or "").strip()
            if not address or not work_type:
                raise ValueError("address and work_type are required")
            sla_date = _as_date(raw.get("sla_date"))
            start = _as_time(raw.get("time_window_start"))
            end = _as_time(raw.get("time_window_end"))
            if start is not None and end is not None and start >= end:
                raise ValueError("invalid time window")
            latitude = _as_float(raw.get("latitude"))
            longitude = _as_float(raw.get("longitude"))
            if (latitude is None) != (longitude is None):
                raise ValueError("latitude and longitude must be supplied together")
            result.append(
                {
                    "external_id": (
                        str(raw["external_id"]).strip()
                        if raw.get("external_id") not in (None, "")
                        else None
                    ),
                    "address": address,
                    "sla_date": sla_date,
                    "work_type": work_type,
                    "time_window_start": start,
                    "time_window_end": end,
                    "latitude": latitude,
                    "longitude": longitude,
                }
            )
        except (TypeError, ValueError) as error:
            errors.append({"row": number, "error": str(error)})
    return result, errors


def _as_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value))


def _as_time(value: Any) -> time | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.time().replace(tzinfo=None)
    if isinstance(value, time):
        return value.replace(tzinfo=None)
    return time.fromisoformat(str(value))


def _as_float(value: Any) -> float | None:
    return None if value in (None, "") else float(value)
